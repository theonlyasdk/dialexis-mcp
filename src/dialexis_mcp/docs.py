"""Core document logic: markdown-IR in, neat files out.

Design for token efficiency:
- Markdown is the only authoring format agents need to emit.
- All create_* functions return a disk path + short summary, never file bytes.
- All read_* functions return markdown with pagination (page/page_size by lines).
"""
from __future__ import annotations

import os
import re
import uuid
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_FORMATS = ("md", "docx", "pptx", "xlsx", "pdf", "html")

_FILENAME_BAD = re.compile(r"[^A-Za-z0-9._-]+")


def get_output_dir() -> Path:
    raw = os.environ.get("DIALEXIS_OUTPUT_DIR", "").strip()
    base = Path(raw).expanduser() if raw else Path.cwd() / "exports"
    base.mkdir(parents=True, exist_ok=True)
    return base.resolve()


def safe_output_path(file_name: str | None, ext: str, title_hint: str = "document") -> Path:
    out = get_output_dir()
    ext = ext.lower().lstrip(".")
    if ext not in SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported format: {ext}. Choose from {SUPPORTED_FORMATS}")
    stem = (file_name or title_hint or "document").strip() or "document"
    # allow subpaths? No — flatten to a safe stem for production safety.
    stem = Path(stem).name
    stem = _FILENAME_BAD.sub("-", stem).strip("-_.") or "document"
    stem = stem[:80]
    candidate = out / f"{stem}.{ext}"
    if candidate.exists():
        candidate = out / f"{stem}-{uuid.uuid4().hex[:8]}.{ext}"
    return candidate


def resolve_read_path(path_str: str) -> Path:
    p = Path(path_str).expanduser()
    if not p.is_absolute():
        # resolve relative to CWD and output dir
        cwd_candidate = (Path.cwd() / p)
        if cwd_candidate.exists():
            p = cwd_candidate
        else:
            out_candidate = get_output_dir() / p.name
            if out_candidate.exists():
                p = out_candidate
    if not p.exists():
        raise FileNotFoundError(f"File not found: {path_str}")
    return p.resolve()


# ---------------------------------------------------------------- markdown IR

@dataclass
class MdBlock:
    kind: str  # heading, paragraph, bullets, numbered, table, code, hr
    level: int = 0
    text: str = ""
    items: list[str] | None = None
    headers: list[str] | None = None
    rows: list[list[str]] | None = None
    language: str = ""


def parse_markdown(md: str) -> list[MdBlock]:
    """Small, dependency-free markdown parser covering what agents emit."""
    lines = (md or "").replace("\r\n", "\n").split("\n")
    blocks: list[MdBlock] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            lang = s[3:].strip()
            buf: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            blocks.append(MdBlock(kind="code", text="\n".join(buf), language=lang))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(MdBlock(kind="heading", level=len(m.group(1)), text=m.group(2).strip()))
            i += 1
            continue
        if re.match(r"^([-*_]\s*){3,}$", s.replace(" ", "")) or re.match(r"^---+$", s):
            blocks.append(MdBlock(kind="hr"))
            i += 1
            continue
        if re.match(r"^(\d+[.)]\s+)", s):
            items: list[str] = []
            while i < n and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(MdBlock(kind="numbered", items=items))
            continue
        if re.match(r"^([-*+]\s+)", s):
            items = []
            while i < n and re.match(r"^\s*[-*+]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*+]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(MdBlock(kind="bullets", items=items))
            continue
        if "|" in s and i + 1 < n and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i + 1]):
            header_cells = [c.strip() for c in s.strip().strip("|").split("|")]
            i += 2
            rows: list[list[str]] = []
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(MdBlock(kind="table", headers=header_cells, rows=rows))
            continue
        # paragraph: gather until blank line or block start
        buf = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not re.match(
            r"^(#{1,6}\s+|```|(\s*[-*+]\s+)|(\s*\d+[.)]\s+)|(\s*\|))", lines[i]
        ):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(MdBlock(kind="paragraph", text=" ".join(b for b in buf if b)))
    return blocks


def extract_title(md: str, fallback: str = "document") -> str:
    for line in (md or "").splitlines():
        m = re.match(r"^#\s+(.*)$", line.strip())
        if m and m.group(1).strip():
            return m.group(1).strip()[:80]
    for line in (md or "").splitlines():
        if line.strip():
            return line.strip()[:80]
    return fallback


def strip_md_inline(text: str) -> str:
    t = text or ""
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"__(.+?)__", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    t = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", t)
    return t


def add_md_runs(paragraph, text: str) -> None:
    """Add runs preserving **bold**, *italic*, `code` for python-docx/pptx."""
    from docx.shared import Pt  # local import so read paths stay light

    token = re.compile(r"(\*\*.+?\*\*|__.+?__|\*.+?\*|_.+?_ |`.+?`)".replace(" ", ""))
    pos = 0
    for m in token.finditer(text or ""):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        chunk = m.group(0)
        if chunk.startswith("**") or chunk.startswith("__"):
            r = paragraph.add_run(chunk[2:-2])
            r.bold = True
        elif chunk.startswith("`"):
            r = paragraph.add_run(chunk[1:-1])
            try:
                r.font.size = Pt(9)
            except Exception:
                pass
        else:
            r = paragraph.add_run(chunk[1:-1])
            r.italic = True
        pos = m.end()
    if pos < len(text or ""):
        paragraph.add_run((text or "")[pos:])


# ---------------------------------------------------------------- writers

def write_docx(blocks: list[MdBlock], title: str, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    # Narrow, readable defaults
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    if title:
        h = doc.add_heading(title, level=0)
        for r in h.runs:
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    for b in blocks:
        if b.kind == "heading" and not (b.level == 1 and b.text == title):
            lvl = min(max(b.level, 1), 4)
            doc.add_heading(strip_md_inline(b.text), level=lvl)
        elif b.kind == "paragraph":
            p = doc.add_paragraph()
            add_md_runs(p, b.text)
        elif b.kind == "bullets":
            for item in b.items or []:
                p = doc.add_paragraph(style="List Bullet")
                add_md_runs(p, item)
        elif b.kind == "numbered":
            for item in b.items or []:
                p = doc.add_paragraph(style="List Number")
                add_md_runs(p, item)
        elif b.kind == "table":
            headers = b.headers or []
            rows = b.rows or []
            table = doc.add_table(rows=1 + len(rows), cols=max(1, len(headers)))
            table.style = "Light Grid Accent 1"
            for j, h in enumerate(headers):
                table.cell(0, j).text = strip_md_inline(h)
            for r_i, row in enumerate(rows, start=1):
                for j in range(len(headers)):
                    table.cell(r_i, j).text = strip_md_inline(row[j] if j < len(row) else "")
            doc.add_paragraph()
        elif b.kind == "code":
            p = doc.add_paragraph()
            r = p.add_run(b.text)
            r.font.size = Pt(9)
            try:
                r.font.name = "Consolas"
            except Exception:
                pass
    doc.save(str(path))


def _split_slides(blocks: list[MdBlock], title: str) -> list[tuple[str, list[MdBlock]]]:
    slides: list[tuple[str, list[MdBlock]]] = []
    cur_title = title
    cur: list[MdBlock] = []
    started = False
    for b in blocks:
        if b.kind == "heading" and b.level <= 2:
            if started:
                slides.append((cur_title, cur))
                cur = []
            cur_title = strip_md_inline(b.text)
            started = True
        else:
            if not started:
                started = True
            cur.append(b)
    if cur or not slides:
        slides.append((cur_title, cur))
    return slides[:60]  # hard cap: token + file sanity


def write_pptx(blocks: list[MdBlock], title: str, path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt as PPt

    prs = Presentation()
    prs.slide_width, prs.slide_height = int(13.333 * 914400), int(7.5 * 914400)
    blank_title: tuple[str, list[MdBlock]] | None = None
    slides = _split_slides(blocks, title or "Presentation")
    if not slides:
        slides = [(title or "Presentation", [])]
    for s_title, s_blocks in slides:
        layout = prs.slide_layouts[1]  # title + content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s_title[:120]
        body = slide.placeholders[1].text_frame
        body.clear()
        first = True
        for b in s_blocks:
            if b.kind in ("bullets", "numbered"):
                for item in (b.items or [])[:12]:
                    p = body.paragraphs[0] if first else body.add_paragraph()
                    p.text = strip_md_inline(item)[:220]
                    p.level = 0
                    first = False
            elif b.kind == "table":
                # tables as bulleted lines in pptx (keeps token cost flat)
                headers = b.headers or []
                if headers:
                    p = body.paragraphs[0] if first else body.add_paragraph()
                    p.text = " | ".join(strip_md_inline(h) for h in headers)[:220]
                    first = False
                for row in (b.rows or [])[:8]:
                    p = body.add_paragraph()
                    p.text = " | ".join(strip_md_inline(c) for c in row)[:220]
                    p.level = 0
            elif b.kind in ("paragraph", "heading"):
                txt = strip_md_inline(b.text)[:220]
                if not txt:
                    continue
                p = body.paragraphs[0] if first else body.add_paragraph()
                p.text = txt
                first = False
            elif b.kind == "code":
                txt = (b.text or "")[:400]
                if txt:
                    p = body.paragraphs[0] if first else body.add_paragraph()
                    p.text = txt
                    first = False
        # readable font size
        for shape in slide.placeholders:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        try:
                            run.font.size = PPt(18 if para == shape.text_frame.paragraphs[0] else 14)
                        except Exception:
                            pass
    prs.save(str(path))


def write_pdf(blocks: list[MdBlock], title: str, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from reportlab.lib import colors
    import html as _html

    def esc(t: str) -> str:
        return _html.escape(strip_md_inline(t or ""))

    styles = getSampleStyleSheet()
    story = []
    if title:
        story += [Paragraph(esc(title), styles["Title"]), Spacer(1, 12)]
    for b in blocks:
        if b.kind == "heading":
            style = styles["Heading1"] if b.level <= 1 else styles["Heading2"] if b.level == 2 else styles["Heading3"]
            story += [Paragraph(esc(b.text), style), Spacer(1, 6)]
        elif b.kind == "paragraph":
            story += [Paragraph(esc(b.text), styles["Normal"]), Spacer(1, 6)]
        elif b.kind in ("bullets", "numbered"):
            for k, item in enumerate(b.items or [], 1):
                prefix = "• " if b.kind == "bullets" else f"{k}. "
                story.append(Paragraph(prefix + esc(item), styles["Normal"]))
            story.append(Spacer(1, 6))
        elif b.kind == "table":
            data = [[esc(h) for h in (b.headers or [])]] + [[esc(c) for c in r] for r in (b.rows or [])]
            if data and data[0]:
                t = Table(data, repeatRows=1)
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
                story += [t, Spacer(1, 6)]
        elif b.kind == "code":
            story += [Paragraph(f"<font face='Courier' size=8>{esc(b.text)}</font>", styles["Code"]), Spacer(1, 6)]
        elif b.kind == "hr":
            story += [HRFlowable(width="100%"), Spacer(1, 6)]
    SimpleDocTemplate(str(path), pagesize=A4, title=title[:100]).build(story)


def write_xlsx(blocks: list[MdBlock], title: str, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "", (title or "Sheet"))[:31] or "Sheet"
    row = 1
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=14)
    row += 2
    for b in blocks:
        if b.kind == "heading":
            ws.cell(row=row, column=1, value=strip_md_inline(b.text)).font = Font(bold=True, size=12)
            row += 1
        elif b.kind in ("paragraph", "code"):
            ws.cell(row=row, column=1, value=strip_md_inline(b.text)[:30000])
            row += 1
        elif b.kind in ("bullets", "numbered"):
            for k, item in enumerate(b.items or [], 1):
                prefix = "" if b.kind == "bullets" else f"{k}. "
                ws.cell(row=row, column=1, value=(prefix + strip_md_inline(item))[:30000])
                row += 1
        elif b.kind == "table":
            ws.cell(row=row, column=1, value="TABLE: " + " | ".join(b.headers or [])).font = Font(bold=True)
            row += 1
            if b.headers:
                for j, h in enumerate(b.headers, 1):
                    ws.cell(row=row, column=j, value=strip_md_inline(h)).font = Font(bold=True)
                row += 1
                for r in (b.rows or []):
                    for j in range(len(b.headers)):
                        ws.cell(row=row, column=j + 1, value=strip_md_inline(r[j] if j < len(r) else "")[:30000])
                    row += 1
            row += 1
        if row > 9000:
            break
    ws.column_dimensions["A"].width = 100
    wb.save(str(path))


def write_md_html(md: str, path: Path) -> None:
    if path.suffix.lower() == ".html":
        import markdown as _md

        body = _md.markdown(md or "", extensions=["tables", "fenced_code"])
        path.write_text(
            f"<!doctype html><html><head><meta charset='utf-8'>"
            f"<title>{strip_md_inline(extract_title(md))}</title></head>"
            f"<body style='max-width:800px;margin:40px auto;font-family:sans-serif'>{body}</body></html>",
            encoding="utf-8",
        )
    else:
        path.write_text((md or "").rstrip() + "\n", encoding="utf-8")


def create_from_markdown(md: str, fmt: str, file_name: str | None = None) -> Path:
    fmt = (fmt or "docx").lower().lstrip(".")
    if fmt not in SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported format '{fmt}'. Choose from {SUPPORTED_FORMATS}")
    if not (md or "").strip():
        raise ValueError("markdown is empty — provide at least a heading or paragraph")
    if len(md) > 400_000:
        raise ValueError("markdown exceeds 400k chars; split into smaller documents")
    title = extract_title(md)
    path = safe_output_path(file_name or title, fmt, title_hint=title)
    blocks = parse_markdown(md)
    if fmt == "docx":
        write_docx(blocks, title, path)
    elif fmt == "pptx":
        write_pptx(blocks, title, path)
    elif fmt == "pdf":
        write_pdf(blocks, title, path)
    elif fmt == "xlsx":
        write_xlsx(blocks, title, path)
    else:
        write_md_html(md, path)
    return path


# ---------------------------------------------------------------- readers

def _paginate(text: str, page: int, page_size: int) -> tuple[str, dict]:
    lines = (text or "").splitlines()
    page = max(1, int(page or 1))
    page_size = min(max(1, int(page_size or 200)), 2000)
    total = len(lines)
    pages = max(1, (total + page_size - 1) // page_size)
    page = min(page, pages)
    chunk = lines[(page - 1) * page_size : page * page_size]
    meta = {"page": page, "pages": pages, "lines": total, "truncated": page < pages}
    return "\n".join(chunk), meta


def extract_markdown(path: Path) -> str:
    """Extract full markdown without pagination. Internal use for convert/edit."""
    ext = path.suffix.lower().lstrip(".")
    if ext in ("md", "markdown", "txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    elif ext == "docx":
        from docx import Document

        doc = Document(str(path))
        out: list[str] = []
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if not t:
                continue
            style = (p.style.name or "").lower()
            if style.startswith("heading 1") or style.startswith("title"):
                out.append(f"# {t}")
            elif style.startswith("heading 2"):
                out.append(f"## {t}")
            elif style.startswith("heading"):
                out.append(f"### {t}")
            elif "bullet" in style:
                out.append(f"- {t}")
            elif "number" in style:
                out.append(f"1. {t}")
            else:
                out.append(t)
        for table in doc.tables:
            rows = [[c.text.strip() for c in r.cells] for r in table.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("| " + " | ".join("---" for _ in rows[0]) + " |")
                out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        text = "\n\n".join(out)
    elif ext == "pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        out = []
        for k, slide in enumerate(prs.slides, 1):
            out.append(f"## Slide {k}")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    out.append(shape.text.strip())
                if shape.has_table:
                    for r in shape.table.rows:
                        out.append(" | ".join(c.text.strip() for c in r.cells))
            out.append("")
        text = "\n".join(out)
    elif ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        out = []
        for k, pg in enumerate(reader.pages, 1):
            out.append(f"## Page {k}\n\n" + ((pg.extract_text() or "").strip()))
        text = "\n\n".join(out)
    elif ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"# Sheet: {ws.title}")
            for r in ws.iter_rows(values_only=True):
                vals = [(str(v) if v is not None else "").strip() for v in r]
                if any(vals):
                    out.append("| " + " | ".join(vals) + " |")
        text = "\n".join(out)
    elif ext == "html":
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        text = soup.get_text("\n").strip()
    else:
        raise ValueError(f"Cannot read .{ext}; supported: md/docx/pptx/xlsx/pdf/html/txt")
    return text


def read_as_markdown(path: Path, detail_level: str = "full", page: int = 1, page_size: int = 200) -> tuple[str, dict]:
    text = extract_markdown(path)
    if detail_level == "summary":
        return "\n".join((text or "").splitlines()[:40]), {"summary": True, "lines": len((text or "").splitlines())}
    if detail_level == "metadata_only":
        return "", {}
    return _paginate(text, page, page_size)


def file_info(path: Path) -> dict:
    st = path.stat()
    info: dict = {"path": str(path), "name": path.name, "format": path.suffix.lstrip(".").lower(), "bytes": st.st_size}
    try:
        ext = path.suffix.lower()
        if ext == ".docx":
            from docx import Document

            doc = Document(str(path))
            info.update({"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)})
        elif ext == ".pptx":
            from pptx import Presentation

            info.update({"slides": len(Presentation(str(path)).slides)})
        elif ext == ".pdf":
            from pypdf import PdfReader

            info.update({"pages": len(PdfReader(str(path)).pages)})
        elif ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True)
            info.update({"sheets": wb.sheetnames})
    except Exception as e:  # metadata must never fail validation
        info["metadata_warning"] = str(e)[:200]
    return info
